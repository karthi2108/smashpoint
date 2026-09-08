import csv
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from ..auth import require_admin
from ..config import CATEGORIES
from ..database import get_db
from ..models import Entry
from ..schemas import EntriesIn, EntryOut

router = APIRouter(prefix="/api/entries", tags=["entries"])


def _check_cat(cat: str) -> str:
    if cat not in CATEGORIES:
        raise HTTPException(status_code=400, detail=f"Unknown category '{cat}'")
    return cat


@router.get("", response_model=list[EntryOut])
def list_entries(category: str | None = None, db: Session = Depends(get_db)):
    q = db.query(Entry)
    if category:
        q = q.filter(Entry.category == _check_cat(category))
    return q.order_by(Entry.id).all()


@router.post("", response_model=list[EntryOut], dependencies=[Depends(require_admin)])
def add_entries(body: EntriesIn, db: Session = Depends(get_db)):
    _check_cat(body.category)
    created = []
    for name in body.names:
        name = name.strip()
        if not name:
            continue
        e = Entry(category=body.category, name=name[:200])
        db.add(e)
        created.append(e)
    db.commit()
    for e in created:
        db.refresh(e)
    return created


@router.post("/import", response_model=list[EntryOut], dependencies=[Depends(require_admin)])
async def import_file(
    category: str = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Import names from the first column of an .xlsx or .csv file."""
    _check_cat(category)
    raw = await file.read()
    names: list[str] = []
    fname = (file.filename or "").lower()

    if fname.endswith((".xlsx", ".xls")):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    names.append(str(row[0]).strip())
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Could not read Excel file: {exc}")
    else:
        try:
            text = raw.decode("utf-8-sig", errors="replace")
            for row in csv.reader(io.StringIO(text)):
                if row and row[0].strip():
                    names.append(row[0].strip())
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Could not read CSV file: {exc}")

    names = [n for n in names if n and n.lower() != "name" and not n.lower().startswith("pair")]
    if not names:
        raise HTTPException(status_code=400, detail="No names found in the first column")

    created = [Entry(category=category, name=n[:200]) for n in names]
    db.add_all(created)
    db.commit()
    for e in created:
        db.refresh(e)
    return created


@router.delete("/{entry_id}", dependencies=[Depends(require_admin)])
def delete_entry(entry_id: int, db: Session = Depends(get_db)):
    e = db.get(Entry, entry_id)
    if not e:
        raise HTTPException(status_code=404, detail="Entry not found")
    db.delete(e)
    db.commit()
    return {"ok": True}


@router.delete("", dependencies=[Depends(require_admin)])
def clear_category(category: str, db: Session = Depends(get_db)):
    _check_cat(category)
    db.query(Entry).filter(Entry.category == category).delete()
    db.commit()
    return {"ok": True}
